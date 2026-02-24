# scripts/convert_empi.py
import os
import re
import html as _html
import openpyxl
from lxml import etree as ET  # używamy lxml (obsługuje CDATA)
from convert import convert_file, INPUT_DIR, OUTPUT_DIR  # główny konwerter

# --------- USTAWIENIA ---------
BRAND_LINKS = {
    "dell":   "https://kompre.pl/pl/c/Laptopy-Dell/364",
    "lenovo": "https://kompre.pl/pl/c/Laptopy-Lenovo/366",
    "hp":     "https://kompre.pl/pl/c/Laptopy-HP/365",
    "apple":  "https://kompre.pl/pl/c/Laptopy-Apple/367",
    "fujitsu":"https://kompre.pl/pl/c/Laptopy-Fujitsu/368",
}
FOOTER_MARK = "<!---->"
LINKS_AS_PLAIN_TEXT = True

# kolumny, których NIE dublujemy w <attrs> (bo już idą w XML jako cat/name/itd.)
EXCLUDED_COLS = {
    "ID oferty",
    "Tytuł oferty",
    "Cena PL",
    "Link do oferty",
    "Status oferty",
    "Liczba sztuk",
    "Kategoria główna",
    "Podkategoria",
    "Zdjęcia",
    "Opis oferty",
}

# --------- POMOCNICZE – Excel: Podkategoria + kolumny od AQ ---------
def _load_excel_maps(in_path: str):
    """
    Czyta Excela i zwraca:
      - extras_by_id: {ID oferty: {nagłówek: wartość, ...}} dla kolumn od AQ w prawo
      - subcat_by_id: {ID oferty: Podkategoria}
    Arkusz: "Szablon" lub pierwszy.
    Nagłówki: wiersz 4, dane: od wiersza 5.
    """
    try:
        wb = openpyxl.load_workbook(in_path, data_only=True, read_only=True)
    except Exception:
        return {}, {}

    ws = wb["Szablon"] if "Szablon" in wb.sheetnames else wb.worksheets[0]

    # nagłówki (wiersz 4)
    headers = [str(c.value).strip() if c.value else "" for c in ws[4]]

    # indeks kolumny ID oferty
    try:
        id_idx = headers.index("ID oferty")
    except ValueError:
        wb.close()
        return {}, {}

    # indeks kolumny Podkategoria (może nie być)
    try:
        subcat_idx = headers.index("Podkategoria")
    except ValueError:
        subcat_idx = None

    # znajdź indeks (0-based) pierwszej kolumny o literze "AQ"
    start_idx = None
    for cell in ws[4]:
        if cell.column_letter == "AQ":
            start_idx = cell.column - 1  # column jest 1-based
            break

    extras_by_id = {}
    subcat_by_id = {}

    if start_idx is None:
        # nie znaleziono AQ – ale nadal możemy mieć Podkategorię
        start_idx = len(headers) + 1  # żeby pętla po extras nic nie dodała

    for row in ws.iter_rows(min_row=5, values_only=True):
        if id_idx >= len(row):
            continue
        raw_id = row[id_idx]
        if raw_id is None:
            continue

        rid = str(raw_id).strip()
        if not rid:
            continue

        # Podkategoria (jeśli jest kolumna)
        if subcat_idx is not None and subcat_idx < len(row):
            sc = row[subcat_idx]
            if sc is not None:
                sc_str = str(sc).strip()
                if sc_str:
                    subcat_by_id[rid] = sc_str

        # dodatkowe kolumny od AQ
        extra = {}
        for i in range(start_idx, len(headers)):
            h = headers[i] if i < len(headers) else ""
            if not h:
                continue
            if h in EXCLUDED_COLS:
                continue
            if i >= len(row):
                continue
            v = row[i]
            if v is None:
                continue
            v_str = str(v).strip()
            if not v_str:
                continue
            extra[h] = v_str

        if extra:
            extras_by_id[rid] = extra

    wb.close()
    return extras_by_id, subcat_by_id


def _enrich_attrs_with_excel(o_el: ET.Element, extras_by_id: dict):
    """
    Dla danego <o> dopisuje wszystkie kolumny z extras_by_id[ID oferty] jako atrybuty <a>,
    jeśli jeszcze nie istnieją w <attrs>.
    """
    oid = (o_el.get("id") or "").strip()
    if not oid:
        return

    row = extras_by_id.get(oid)
    if not row:
        return

    attrs_el = o_el.find("attrs")
    if attrs_el is None:
        attrs_el = ET.SubElement(o_el, "attrs")

    existing = {(a.get("name") or "").strip() for a in attrs_el.findall("a")}

    for col_name, val in row.items():
        if not col_name or not val:
            continue
        if col_name in existing:
            continue
        ET.SubElement(attrs_el, "a", {"name": col_name}).text = val

# --------- POMOCNICZE – stopka / opis ---------
def _collect_attrs(o_el):
    out = {}
    attrs_el = o_el.find("attrs")
    if attrs_el is None:
        return out
    for a in attrs_el.findall("a"):
        name = (a.get("name") or "").strip()
        val = (a.text or "").strip()
        if name:
            out[name] = val
    return out

def _brand(attrs):
    return (attrs.get("Producent") or "").strip()

def _warranty(attrs):
    gw = (attrs.get("Informacje o gwarancjach") or attrs.get("Gwarancja") or "").strip()
    return gw

def _category(o_el):
    cat_el = o_el.find("cat")
    return (cat_el.text or "").strip() if cat_el is not None else ""

def _name(o_el):
    name_el = o_el.find("name")
    return (name_el.text or "").strip() if name_el is not None else ""

def _build_link_block(kategoria, producent):
    if not kategoria or not producent:
        return ""
    kat = kategoria.lower()
    brand = producent.lower()
    if "laptop" in kat:
        url = BRAND_LINKS.get(brand)
    elif "komputer" in kat:
        url = "https://kompre.pl/pl/c/Komputery-Stacjonarne/34"
    elif "monitor" in kat:
        url = "https://kompre.pl/monitory"
    else:
        return ""
    if not url:
        return ""
    return (
        f"<p>Posiadamy też inne modele {producent} – sprawdź: {url}. "
        f"Każdy egzemplarz jest testowany, czyszczony i przygotowany do pracy z aktualnym systemem. "
        f"Długa gwarancja door-to-door zapewnia wsparcie i bezpieczeństwo zakupu.</p>"
    )

def _build_footer_html(name, producent, gwarancja, kategoria):
    link_block = _build_link_block(kategoria, producent)
    gw = (gwarancja or "12 miesięcy").strip()
    _ = gw  # na razie niewykorzystane w tekście
    return (
        f'{FOOTER_MARK}'
        f'<hr/><p><strong>{name}</strong> pochodzi z oferty <strong>Kompre.pl</strong> – '
        f'autoryzowanego sprzedawcy komputerów poleasingowych klasy biznes.</p> '
        f'{link_block}'
    )

def _inner_html(el: ET.Element) -> str:
    parts = []
    if el.text:
        parts.append(el.text)
    for c in el:
        parts.append(ET.tostring(c, encoding="unicode"))
    return "".join(parts)

def _set_desc_cdata(desc_el: ET.Element, html_string: str):
    desc_el.clear()
    desc_el.text = ET.CDATA(html_string)

def _already_has_footer(html: str) -> bool:
    return (FOOTER_MARK in html) or ("Kompre.pl" in html and "door-to-door" in html)

def _append_footer_to_desc(o_el):
    desc_el = o_el.find("desc")
    if desc_el is None:
        return
    current_html = _inner_html(desc_el)
    if _already_has_footer(current_html):
        return
    attrs = _collect_attrs(o_el)
    name = _name(o_el)
    producent = _brand(attrs)
    gwarancja = _warranty(attrs)
    kategoria = _category(o_el)
    footer_html = _build_footer_html(name, producent, gwarancja, kategoria)
    joiner = "\n" if current_html and not current_html.endswith("\n") else ""
    new_html = f"{current_html}{joiner}{footer_html}".strip()
    _set_desc_cdata(desc_el, new_html)

# --- Sanizacja HTML opisu ---
_SCRIPT_IFRAME_IMG_RE = re.compile(
    r"(?is)<script.*?</script>|<iframe.*?</iframe>|<img\b[^>]*>"
)

def _sanitize_basic(html_str: str) -> str:
    return _SCRIPT_IFRAME_IMG_RE.sub("", html_str or "")

def _has_html_tags(s: str) -> bool:
    return bool(re.search(r"<[a-zA-Z][^>]*>", s or ""))

# --- Edycje copy w opisie (reguły) ---
def _apply_copy_edits(s: str) -> str:
    rules = [
        (re.compile(r'(?i)Nawiąż kontakt z kim tylko chcesz'), 'Nawiąż znajomość z kim tylko chcesz'),
        (re.compile(r'(?i)Świetny stosunek jakości do ceny'), 'Świetna jakość'),
        (re.compile(r'(?i)\bw\s+gratisie\b'), ''),     # usuń "w Gratisie"
        (re.compile(r'(?i)\bgratis!?\b'), ''),        # usuń "Gratis" / "GRATIS!"
        (re.compile(r'(?i)Nie tylko cena,\s*'), ''),  # usuń "Nie tylko cena,"
        (re.compile(r'(?i)\bcenie\b'), 'ofercie'),
        (re.compile(r'(?i)\bcena\b'), 'ofercie'),
        (re.compile(r'(?i)Kup teraz'), ''),           # usuń "Kup teraz"
    ]
    out = s
    for rx, repl in rules:
        out = rx.sub(repl, out)
    out = re.sub(r'\s{2,}', ' ', out)
    return out.strip()

def _normalize_html_structure(html_str: str) -> str:
    if not html_str:
        return ""

    s = html_str

    # 1) Usuń nagłówki-separatory typu ________
    s = re.sub(r'<h[1-6]>_+</h[1-6]>', '', s, flags=re.IGNORECASE)

    # 2) Zamień wszystkie H1–H6 na H2
    s = re.sub(
        r'<h[1-6]>(.*?)</h[1-6]>',
        r'<h2>\1</h2>',
        s,
        flags=re.IGNORECASE | re.DOTALL,
    )

    # 3) Scal sąsiadujące listy <ul>...</ul><ul>...</ul> -> jedna lista
    s = re.sub(r'</ul>\s*<ul>', '', s, flags=re.IGNORECASE)

    # 4) Wywal puste paragrafy
    s = re.sub(
        r'<p>(?:\s|&nbsp;|<br\s*/?>)*</p>',
        '',
        s,
        flags=re.IGNORECASE,
    )

    # 5) Zbij nadmiar białych znaków między tagami
    s = re.sub(r'>\s+<', '><', s)

    return s.strip()

def _force_desc_cdata(o_el: ET.Element):
    """Opis w realnym HTML (CDATA), bez <img>, z poprawkami copy + normalizacją struktury."""
    desc_el = o_el.find("desc")
    if desc_el is None:
        return
    raw = _inner_html(desc_el).strip()
    unescaped = _html.unescape(raw).strip()
    cleaned = _sanitize_basic(unescaped)
    cleaned = _apply_copy_edits(cleaned)
    cleaned = _normalize_html_structure(cleaned)
    if not _has_html_tags(cleaned) and cleaned:
        cleaned = f"<p>{cleaned}</p>"
    _set_desc_cdata(desc_el, cleaned)

# --- Formatowanie pojemności ---
def _format_capacity_unit(val: str) -> str:
    if not val:
        return ""
    m = re.search(r"(\d+(?:[.,]\d+)?)", val)
    if not m:
        return ""
    num = m.group(1).replace(",", ".")
    try:
        f = float(num)
    except Exception:
        return ""
    i = int(round(f))
    return "1 TB" if i == 1 else f"{i} GB"

# --- Normalizacja cali w 'Przekątna ekranu' ---
def _normalize_inches(value: str) -> str:
    """Zwraca N[.N]\" (np. 14\", 12\"). Usuwa 'cali' itp., dokleja jeśli brak."""
    if not value:
        return value
    m = re.search(r"(\d+(?:[.,]\d+)?)", value)
    if not m:
        v = value.strip()
        return v if v.endswith('"') else (v + '"')
    num = m.group(1).replace(",", ".")
    if "." in num:
        num = num.rstrip("0").rstrip(".")
    return f'{num}"'

# --------- GŁÓWNA LOGIKA KONWERSJI ---------
def convert_file_empi(in_path, out_path):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    temp_path = os.path.join(OUTPUT_DIR, "_temp_base.xml")

    # bazowy XML z convert.py (Morele / inne feedy dalej używają convert.py normalnie)
    convert_file(in_path, temp_path)

    # Excel: dodatkowe kolumny od AQ + Podkategoria
    extras_by_id, subcat_by_id = _load_excel_maps(in_path)

    parser = ET.XMLParser(remove_blank_text=True)
    tree = ET.parse(temp_path, parser)
    root = tree.getroot()

    for o in root.findall("o"):
        oid = (o.get("id") or "").strip()

        # najpierw dołóż dodatkowe atrybuty z Excela (od AQ)
        _enrich_attrs_with_excel(o, extras_by_id)

        # Podkategoria → <subcat>
        if oid and oid in subcat_by_id:
            sub_val = subcat_by_id[oid]
            sub_el = o.find("subcat")
            if sub_el is None:
                ET.SubElement(o, "subcat").text = sub_val
            else:
                sub_el.text = sub_val

        # dostępność: aktywna tylko gdy stock >= 4
        try:
            stock_num = int(o.get("stock", "0"))
        except Exception:
            try:
                stock_num = int(float(o.get("stock", "0")))
            except Exception:
                stock_num = 0
        if o.get("avail") == "1" and stock_num < 4:
            o.set("avail", "99")
            o.set("stock", "0")
            o.set("basket", "0")

        # dopisz "poleasingowe" do kategorii
        cat_el = o.find("cat")
        if cat_el is not None and cat_el.text:
            cat_text = cat_el.text.strip()
            norm = cat_text.lower()
            if "poleasingowe" not in norm:
                if norm == "laptopy":
                    cat_el.text = "Laptopy poleasingowe"
                elif norm == "komputery":
                    cat_el.text = "Komputery poleasingowe"
                elif norm == "monitory komputerowe":
                    cat_el.text = "Monitory poleasingowe"

        # usuń desc_json (empi korzysta z HTML)
        for dj in o.findall("desc_json"):
            parent = dj.getparent() if hasattr(dj, "getparent") else o
            parent.remove(dj)

        # --- ATRYBUTY: transformacje dla empi ---
        attrs_el = o.find("attrs")
        if attrs_el is not None:
            # słownik atrybutów
            attrs = {}
            for a in attrs_el.findall("a"):
                name = (a.get("name") or "").strip()
                val = (a.text or "").strip()
                if name:
                    attrs[name] = val

            # 1) Stan: Używany -> Poleasingowy
            for a in attrs_el.findall("a"):
                if (a.get("name") or "").strip() == "Stan":
                    val = (a.text or "").strip()
                    if re.search(r"\bużywany\b", val, flags=re.IGNORECASE):
                        a.text = "Poleasingowy"

            # 2) Zmiany nazw atrybutów RAM / ekran / rozdzielczość
            for a in attrs_el.findall("a"):
                n = (a.get("name") or "").strip()
                if n == 'Wielkość pamięci RAM':
                    a.set("name", "Pamięć RAM (zainstalowana)")
                elif n == 'Przekątna ekranu ["]':
                    a.set("name", "Przekątna ekranu")
                elif n == 'Rozdzielczość (px)':
                    a.set("name", "Rozdzielczość")

            # 2a) Przekątna ekranu – wymuś format N[.N]"
            for a in attrs_el.findall("a"):
                if (a.get("name") or "").strip() == "Przekątna ekranu":
                    v = (a.text or "").strip()
                    if v:
                        a.text = _normalize_inches(v)

            # 3) Ekran dotykowy: tylko "Nie" lub "z ekranem dotykowym"
            for a in attrs_el.findall("a"):
                if (a.get("name") or "").strip() == "Ekran dotykowy":
                    v = (a.text or "").strip().lower()
                    if v == "tak":
                        a.text = "z ekranem dotykowym"
                    elif v == "nie":
                        a.text = "Nie"

            # 4) Dyski SSD/HDD z pojemnością
            typ = (attrs.get("Typ dysku twardego") or "").lower()
            cap_raw = attrs.get("Pojemność dysku [GB]") or ""
            cap_fmt = _format_capacity_unit(cap_raw)
            if cap_fmt:
                if "ssd" in typ and not any((x.get("name") or "") == "Dysk SSD" for x in attrs_el.findall("a")):
                    ET.SubElement(attrs_el, "a", {"name": "Dysk SSD"}).text = cap_fmt
                if "hdd" in typ and not any((x.get("name") or "") == "Dysk HDD" for x in attrs_el.findall("a")):
                    ET.SubElement(attrs_el, "a", {"name": "Dysk HDD"}).text = cap_fmt

            # 4a) Grafika zintegrowana -> dopisz pamięć karty jako "Współdzielona z RAM"
            rodzaj = (attrs.get("Rodzaj karty graficznej") or "").strip().lower()
            if "zintegrowana" in rodzaj:
                has_mem = any((x.get("name") or "").strip() == "Pamięć karty graficznej" for x in attrs_el.findall("a"))
                if not has_mem:
                    ET.SubElement(attrs_el, "a", {"name": "Pamięć karty graficznej"}).text = "Współdzielona z RAM"

            # 5) Informacje o gwarancjach -> Gwarancja (liczba)
            for a in attrs_el.findall("a"):
                if (a.get("name") or "").strip().lower() == "informacje o gwarancjach":
                    text = (a.text or "").strip()
                    m = re.search(r"(\d+)", text)
                    value = m.group(1) if m else ""
                    a.set("name", "Gwarancja")
                    a.text = value

        # --- OPIS: HTML w CDATA (bez IMG) + poprawki copy + stopka
        _force_desc_cdata(o)
        _append_footer_to_desc(o)

    tree.write(out_path, encoding="utf-8", xml_declaration=True, pretty_print=True)
    try:
        os.remove(temp_path)
    except FileNotFoundError:
        pass

    print(f"[empi OK] Zapisano: {out_path}")

def main():
    for name in os.listdir(INPUT_DIR):
        if name.lower().endswith((".xlsm", ".xlsx", ".xls")):
            src = os.path.join(INPUT_DIR, name)
            dst = os.path.join(OUTPUT_DIR, "empi.xml")
            print(f"[empi] {src} -> {dst}")
            convert_file_empi(src, dst)
            break

if __name__ == "__main__":
    main()
